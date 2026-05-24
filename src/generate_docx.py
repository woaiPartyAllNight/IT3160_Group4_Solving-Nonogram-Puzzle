import os
import sys
import subprocess

# Tự động cài đặt thư viện python-docx nếu chưa có
try:
    import docx
except ImportError:
    print("[*] Library 'python-docx' not found. Installing automatically...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = Document()
    
    # Thiết lập lề trang (2.54 cm cho tất cả các bên)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Cấu hình Style mặc định
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(12)
    font.color.rgb = RGBColor(51, 51, 51) # Màu xám đen lịch sự

    # Hàm tạo Header trang trí
    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(26, 82, 118) # Màu xanh dương đậm
        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Hàm thêm Tiêu đề mục lớn (Heading 1)
    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(33, 97, 140)
        
        # Thêm đường viền dưới tiêu đề lớn để trang trí
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '12') # Độ dày nét vẽ
        bottom.set(qn('w:space'), '4')
        bottom.set(qn('w:color'), '5DADE2') # Màu xanh nhạt
        pbdr.append(bottom)
        pPr.append(pbdr)

    # Hàm thêm Tiêu đề mục nhỏ (Heading 2)
    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(40, 116, 166)

    # Hàm thêm văn bản thông thường
    def add_body(text, bold_prefix="", italic_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(6)
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
        if italic_prefix:
            r_italic = p.add_run(italic_prefix)
            r_italic.italic = True
        p.add_run(text)
        return p

    # Hàm thêm chú thích dạng khối hộp (Callout box) để tạo điểm nhấn
    def add_callout(text, title="Lưu ý quan trọng"):
        table = doc.add_table(rows=1, cols=1)
        table.autofit = False
        table.columns[0].width = Inches(6.5)
        cell = table.cell(0, 0)
        
        # Đổ màu nền nhạt
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="F2F4F4"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))
        
        # Tạo đường viền dày màu xanh ở bên trái
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        left = OxmlElement('w:left')
        left.set(qn('w:val'), 'single')
        left.set(qn('w:sz'), '24') # Rộng 3pt
        left.set(qn('w:color'), '3498DB')
        tcBorders.append(left)
        # Ẩn các đường viền còn lại
        for border_name in ['top', 'bottom', 'right']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'none')
            tcBorders.append(border)
        tcPr.append(tcBorders)
        
        # Viết nội dung vào ô
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        run_title = p.add_run(f"★ {title}\n")
        run_title.bold = True
        run_title.font.color.rgb = RGBColor(41, 128, 185)
        
        run_text = p.add_run(text)
        run_text.font.size = Pt(11)
        run_text.font.color.rgb = RGBColor(86, 101, 115)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- BẮT ĐẦU SOẠN THẢO BÁO CÁO ---

    add_title("BÁO CÁO CHI TIẾT: CÁC QUY TẮC LOGIC TRONG THUẬT TOÁN NONOGRAM SOLVER")
    
    add_body(
        "Báo cáo này trình bày chi tiết về cấu trúc dữ liệu giới hạn và 10 quy tắc logic cốt lõi "
        "đang được áp dụng trong thư mục 'solver/rules/'. Các quy tắc này đóng vai trò quan trọng trong việc "
        "thu hẹp không gian tìm kiếm thông qua kỹ thuật Lan truyền ràng buộc (Constraint Propagation), giúp bộ giải "
        "giải quyết các bảng Nonogram kích thước lớn vô cùng nhanh chóng mà không cần thử sai (backtrack) quá nhiều."
    )

    add_heading_1("I. KHÁI NIỆM CỐT LÕI: KHÔNG GIAN GIỚI HẠN (BOUND)")
    add_body(
        "Để thực hiện suy luận logic, hệ thống duy trì một biến trạng thái cực kỳ quan trọng mang tên ",
        bold_prefix="bound"
    )
    add_body(
        "Biến này định nghĩa phạm vi di chuyển khả dĩ của từng khối đen trên mỗi hàng/cột:\n"
        "  • bound[0][i][k] = [earliest_start, latest_end]: Đại diện cho vị trí bắt đầu sớm nhất và vị trí kết thúc muộn nhất có thể của khối k trên hàng i.\n"
        "  • bound[1][j][k] = [earliest_start, latest_end]: Tương tự cho cột j.\n"
        "Mỗi khi một ô đen (1) hoặc ô chéo (-1) được xác định, các quy tắc dưới đây sẽ hoạt động để kéo sớm nhất (earliest_start) tăng lên và ép muộn nhất (latest_end) giảm xuống. "
        "Khi khoảng cách giữa chúng bằng đúng độ dài khối, vị trí của khối đó được cố định hoàn toàn."
    )

    add_heading_1("II. CHI TIẾT 10 QUY TẮC LOGIC ĐANG ÁP DỤNG TRONG CODE")

    # 1. Rule 1.1
    add_heading_2("1. Quy tắc 1.1: Chồng lấp cực hạn (Overlapping Fill)")
    add_body("apply_rule1_1_block_row / apply_rule1_1_block_col", italic_prefix="Hàm thực thi: ")
    add_body(
        "Khi ta biết vị trí bắt đầu sớm nhất (nếu khối được đẩy sát về bên trái) và vị trí kết thúc muộn nhất (nếu khối đẩy sát về bên phải) "
        "của khối k. Nếu chiều dài của khối k đủ lớn để hai trạng thái biên này giao nhau (chồng lấp lên nhau), thì toàn bộ phần giao nhau "
        "đó chắc chắn phải là các ô đen (1).\n"
        "Ví dụ: Hàng dài 10 ô, gợi ý có 1 khối độ dài 8. "
        "Vị trí sớm nhất là ô [0 đến 7]. Vị trí muộn nhất là ô [2 đến 9]. "
        "Phần chồng lấp là ô [2 đến 7] (độ dài 6 ô) chắc chắn là màu đen."
    )

    # 2. Rule 1.2 + 1.4
    add_heading_2("2. Quy tắc 1.2 & 1.4: Loại trừ ô trống quá ngắn (Short Gap Pruning)")
    add_body("Phần logic tích hợp trong apply_rule1_cell_scanning_row / _col", italic_prefix="Hàm thực thi: ")
    add_body(
        "Nếu trên hàng/cột xuất hiện một khoảng trống (vây quanh bởi các dấu X hoặc biên bảng) mà chiều dài của khoảng trống này "
        "nhỏ hơn kích thước của khối nhỏ nhất có thể nằm tại đó, thì không khối nào có thể xếp vào đây. Toàn bộ các ô trống trong khoảng "
        "này sẽ lập tức được đánh dấu X (-1).\n"
        "Ví dụ: Một ô trống đơn độc (dài 1 ô) nằm giữa hai ô X, trong khi tất cả các khối gợi ý còn lại đều có độ dài từ 2 trở lên. "
        "Ô trống đó chắc chắn phải là dấu X (-1)."
    )

    # 3. Rule 1.3
    add_heading_2("3. Quy tắc 1.3: Đánh dấu ô ngăn cách biên (Spacer Placement)")
    add_body("Phần đầu và cuối của apply_rule1_cell_scanning_row / _col", italic_prefix="Hàm thực thi: ")
    add_body(
        "Khi một khối có kích thước bằng 1 nằm ở biên của giới hạn khả dĩ, nếu ô tại biên đã được tô đen (1), "
        "thì ô liền kề ngoài giới hạn chắc chắn phải là dấu X (-1) để đảm bảo tính ngăn cách giữa các khối (Nonogram quy định các khối đen phải cách nhau ít nhất một ô trống).\n"
        "Ví dụ: Khối k có độ dài 1 chỉ có thể nằm trong khoảng ô [3 đến 4]. Nếu ô [3] đã là màu đen, thì ô [4] bắt buộc phải là X."
    )

    # 4. Rule 1.5.1-3
    add_heading_2("4. Quy tắc 1.5.1-3: Lan truyền ô đen (Black Run Extension)")
    add_body("Phần quét ô đen trong apply_rule1_cell_scanning_row / _col", italic_prefix="Hàm thực thi: ")
    add_body(
        "Khi phát hiện một ô đen (1), dựa vào thông tin về độ dài tối thiểu của khối có thể chứa nó (min_run), thuật toán sẽ quét "
        "và tô đen các ô xung quanh nếu việc không tô đen chúng sẽ dẫn đến mâu thuẫn (ví dụ: tạo thành khối quá ngắn hoặc đụng dấu X).\n"
        "Ví dụ: Nếu ô [5] là ô đen và khối khả dĩ duy nhất bao phủ nó có độ dài là 3. Nếu ô [4] là dấu X, thì các ô [6] và [7] bắt buộc phải được tô đen."
    )

    # 5. Rule 1.5.4
    add_heading_2("5. Quy tắc 1.5.4: Chặn hai đầu khối đen hoàn chỉnh (Block Sealing)")
    add_body("Tích hợp trong apply_rule1_cell_scanning_row / _col", italic_prefix="Hàm thực thi: ")
    add_body(
        "Khi ta tìm thấy một phân đoạn ô đen liên tục có chiều dài đúng bằng kích thước khối gợi ý duy nhất có thể thuộc về nó, "
        "khối này đã hoàn thành. Hệ thống sẽ tự động đặt dấu X (-1) ở hai đầu phân đoạn đó để 'đóng hộp' khối này lại.\n"
        "Ví dụ: Một khối gợi ý có độ dài 4. Trên bảng đã xuất hiện chuỗi 4 ô đen liên tiếp. Hai ô ngay trước và ngay sau chuỗi này chắc chắn phải là X (-1)."
    )

    # 6. Rule 2.1
    add_heading_2("6. Quy tắc 2.1: Thắt chặt ranh giới lân cận (Neighbor Bound Tightening)")
    add_body("apply_rule2_1_block_row / apply_rule2_1_block_col", italic_prefix="Hàm thực thi: ")
    add_body(
        "Giới hạn của các khối đen liền kề có sự phụ thuộc tuyến tính chặt chẽ lẫn nhau. Khối k không thể bắt đầu trước điểm kết thúc sớm nhất "
        "của khối k-1 cộng với 1 ô trống ngăn cách. Ngược lại, khối k không thể kết thúc sau điểm bắt đầu muộn nhất của khối k+1 trừ đi 1 ô trống.\n"
        "Công thức: \n"
        "  • bound_sớm_nhất[k] = bound_sớm_nhất[k-1] + độ_dài[k-1] + 1\n"
        "  • bound_muộn_nhất[k] = bound_muộn_nhất[k+1] - độ_dài[k+1] - 1"
    )

    # 7. Rule 2.2
    add_heading_2("7. Quy tắc 2.2: Dịch chuyển giới hạn bằng ô đen biên (Black Border Push)")
    add_body("apply_rule2_2_block_row / apply_rule2_2_block_col", italic_prefix="Hàm thực thi: ")
    add_body(
        "Nếu ô nằm ngay sát trước giới hạn bắt đầu sớm nhất hoặc ngay sát sau giới hạn kết thúc muộn nhất là ô đen (1), "
        "ta buộc phải dịch chuyển giới hạn đó vào trong 1 ô. Vì nếu giữ nguyên giới hạn cũ, ô đen biên này sẽ bị gộp vào khối k "
        "khiến khối k bị dài ra quá kích thước cho phép hoặc vi phạm khoảng cách trống.\n"
        "Quy tắc này thực hiện 4 bước kiểm tra biên (Trực quan hóa bằng các ghi chú ngắn gọn Việt hóa trong code)."
    )

    # 8. Rule 2.3
    add_heading_2("8. Quy tắc 2.3: Đẩy khối nhỏ dựa trên phân đoạn đen (Segment Bound Push)")
    add_body("apply_rule2_3_segment_row / apply_rule2_3_segment_col", italic_prefix="Hàm thực thi: ")
    add_body(
        "Khi phát hiện một phân đoạn ô đen đã tô có độ dài thực tế là 'seg'. Nếu một khối k nào đó có độ dài nhỏ hơn 'seg', "
        "thì khối k này tuyệt đối không thể chứa phân đoạn đen đó (vì nếu chứa sẽ làm khối bị dài quá quy định). "
        "Do đó, giới hạn khả dĩ của các khối nhỏ hơn sẽ bị đẩy ra xa khỏi phân đoạn này (quét xuôi và quét ngược).\n"
        "Ví dụ: Phân đoạn đen dài 4 ô. Khối k có độ dài gợi ý chỉ là 2. Giới hạn di chuyển của khối k sẽ bị thu hẹp lại để không bao giờ chạm vào phân đoạn đen 4 ô kia."
    )

    # 9. Rule 3.1
    add_heading_2("9. Quy tắc 3.1: Kết nối neo đen trong cùng khối (Black Anchor Connection)")
    add_body("apply_rule3_1_block_row / apply_rule3_1_block_col", italic_prefix="Hàm thực thi: ")
    add_body(
        "Nếu trong phạm vi cô lập của khối k xuất hiện hai ô đen riêng biệt được ngăn cách bởi các ô trống (chưa xác định). "
        "Nếu chúng chắc chắn phải thuộc về cùng một khối k, thì toàn bộ các ô nằm ở khoảng giữa hai ô đen này bắt buộc phải được tô đen "
        "để tạo thành một khối liền mạch duy nhất.\n"
        "Ví dụ: Khối k dài 5 ô chỉ có thể nằm trong khoảng ô [2 đến 8]. Hiện tại ô [3] và ô [6] đã đen. Toàn bộ khoảng giữa là ô [4], [5] bắt buộc phải tô đen."
    )

    # 10. Rule 3.2 & 3.3
    add_heading_2("10. Quy tắc 3.2 & 3.3: Lan truyền neo biên và khoảng trống cô lập")
    add_body("apply_rule3_2_block... & apply_rule3_3_first_block... / apply_rule3_3_last_block...", italic_prefix="Hàm thực thi: ")
    add_body(
        "• Quy tắc 3.2: Xử lý và đánh dấu X các khoảng trống dọc/ngang cô lập quá ngắn và thắt chặt giới hạn vào các vùng trống đủ lớn.\n"
        "• Quy tắc 3.3: Dành riêng cho khối đầu tiên hoặc khối cuối cùng của hàng/cột. Nếu ô ở biên giới hạn đã là ô đen (1), "
        "khối đó chắc chắn phải neo cố định tại đây. Thuật toán sẽ lập tức tô đen toàn bộ chiều dài khối, đặt dấu X để chặn đầu còn lại, "
        "và đồng thời dịch chuyển giới hạn của các khối tiếp theo lùi ra xa."
    )

    add_heading_1("III. ĐÁNH GIÁ HIỆU QUẢ CỦA CÁC QUY TẮC LOGIC")
    add_body(
        "Nhờ việc kết hợp đồng thời cả 10 quy tắc trên trong một vòng lặp quét liên tục (Constraint Propagation) ở file 'logical.py', "
        "hầu hết các ô trên bảng Nonogram đều được giải quyết tự động bằng suy luận toán học thuần túy. "
        "Điều này giúp giảm thiểu số lần phải 'thử sai' trong thuật toán Quay lui (Backtracking) xuống gần bằng 0 đối với các bảng dễ/trung bình, "
        "và giảm từ hàng triệu nhánh xuống còn vài chục nhánh đối với các bảng cực khó, giúp phần mềm hoạt động trơn tru và tối ưu nhất."
    )
    
    add_callout(
        "Báo cáo này được cấu trúc bằng thư viện python-docx với định dạng chuyên nghiệp. "
        "Bạn có thể mở trực tiếp file '.docx' này bằng Microsoft Word, Google Docs hoặc LibreOffice để chỉnh sửa nội dung, "
        "căn lề hay thay đổi font chữ một cách dễ dàng.",
        title="Hướng dẫn sử dụng file báo cáo"
    )

    # Lưu tài liệu
    output_path = os.path.join(os.path.dirname(__file__), "Bao_cao_Nonogram_Rules.docx")
    doc.save(output_path)
    print(f"[+] Created report successfully at: {output_path}")

if __name__ == "__main__":
    create_report()
